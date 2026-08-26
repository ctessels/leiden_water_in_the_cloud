"""Maak de dagelijkse zeven-daagse voorspelling met het getrainde sensormodel."""

from __future__ import annotations

import argparse
import json
import shutil
import sqlite3
import tempfile
from contextlib import closing
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt


PROJECT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_DIR / "data_collection" / "sensor_data"
DB_PATH = DATA_DIR / "database.db"
WEATHER_PATH = DATA_DIR / "voorschoten_weerdata.csv"
USABLE_SENSORS_PATH = DATA_DIR / "usable_sensors.json"
ARTIFACT_DIR = Path(__file__).resolve().parent / "model_artifacts"
OUTPUT_DIR = Path(__file__).resolve().parent / "prediction_results"
VISUALISATION_DIR = Path(__file__).resolve().parent / "prediction_visualisation"
MODEL_PATH = ARTIFACT_DIR / "archived_sensor_model.json"

FORECAST_HORIZON_DAYS = 7
MINIMUM_TRAINING_DAYS = 90
RAIN_MEMORY_DAYS = 5
MAX_SENSOR_AGE_DAYS = 2
MAX_FORECAST_AGE_DAYS = 1
DEFAULT_SIMULATIONS = 1_000


def load_usable_sensor_ids(path: Path = USABLE_SENSORS_PATH) -> set[int]:
    """Lees alle sensoren uit alle duurcategorieen."""
    with path.open(encoding="utf-8") as source:
        sensors_by_category = json.load(source)
    return {
        int(sensor_id)
        for sensor_ids in sensors_by_category.values()
        for sensor_id in sensor_ids
    }


def load_active_probe_inventory(db_path: Path = DB_PATH) -> pd.DataFrame:
    """Maak een regel voor iedere probe van iedere actieve sensor."""
    with closing(sqlite3.connect(db_path)) as connection:
        sensors = pd.read_sql_query(
            """
            SELECT
                device_id,
                location_name,
                measuring_points,
                mp_depth_1,
                mp_depth_2,
                mp_depth_3,
                usable_from
            FROM DimSensor
            WHERE is_active = 1
            ORDER BY device_id
            """,
            connection,
        )

    rows = []
    for sensor in sensors.itertuples(index=False):
        depths = {1: sensor.mp_depth_1, 2: sensor.mp_depth_2, 3: sensor.mp_depth_3}
        for probe_number in range(1, int(sensor.measuring_points or 0) + 1):
            rows.append(
                {
                    "device_id": int(sensor.device_id),
                    "location_name": sensor.location_name,
                    "probe_number": probe_number,
                    "depth_cm": depths.get(probe_number),
                    "usable_from": sensor.usable_from,
                }
            )
    return pd.DataFrame(rows)


def load_daily_sensor_data(db_path: Path = DB_PATH) -> pd.DataFrame:
    """Lees alle actieve probes en aggregeer naar lokale kalenderdagen."""
    with closing(sqlite3.connect(db_path)) as connection:
        sensor_data = pd.read_sql_query(
            """
            SELECT
                fsd.device_id,
                ds.location_name,
                fsd.probe_number,
                CASE fsd.probe_number
                    WHEN 1 THEN ds.mp_depth_1
                    WHEN 2 THEN ds.mp_depth_2
                    WHEN 3 THEN ds.mp_depth_3
                END AS depth_cm,
                ds.usable_from,
                fsd.timestamp,
                fsd.temperature,
                fsd.relative_permittivity
            FROM FactSensorData AS fsd
            INNER JOIN DimSensor AS ds
                ON ds.device_id = fsd.device_id
            WHERE ds.is_active = 1
            ORDER BY fsd.device_id, fsd.probe_number, fsd.timestamp
            """,
            connection,
        )

    sensor_data["timestamp"] = pd.to_datetime(
        sensor_data["timestamp"], unit="s", utc=True, errors="coerce"
    )
    sensor_data["date"] = sensor_data["timestamp"].dt.tz_convert(
        "Europe/Amsterdam"
    ).dt.date
    sensor_data["relative_permittivity"] = pd.to_numeric(
        sensor_data["relative_permittivity"], errors="coerce"
    )
    sensor_data["temperature"] = pd.to_numeric(
        sensor_data["temperature"], errors="coerce"
    )
    sensor_data["usable_from_date"] = pd.to_datetime(
        sensor_data["usable_from"], errors="coerce"
    ).dt.date
    sensor_data = sensor_data[
        sensor_data["usable_from_date"].isna()
        | (sensor_data["date"] >= sensor_data["usable_from_date"])
    ].dropna(subset=["date", "relative_permittivity", "depth_cm"])

    return (
        sensor_data.groupby(
            ["device_id", "location_name", "probe_number", "depth_cm", "date"],
            as_index=False,
            dropna=False,
        )
        .agg(
            relative_permittivity=("relative_permittivity", "mean"),
            ground_temperature=("temperature", "mean"),
            raw_readings=("relative_permittivity", "size"),
        )
        .sort_values(["depth_cm", "device_id", "probe_number", "date"])
        .reset_index(drop=True)
    )


def calculate_energy_proxy(
    sunshine_minutes: pd.Series,
    average_temperature_c: pd.Series,
    maximum_temperature_c: pd.Series,
) -> pd.Series:
    """Bereken de energieproxy uit zonneschijnduur en luchttemperatuur."""
    return sunshine_minutes * (
        (average_temperature_c + maximum_temperature_c) / 2.0
    )


def load_historical_weather(path: Path = WEATHER_PATH) -> pd.DataFrame:
    """Lees KNMI-dagwaarden en zet ze om naar forecast-eenheden."""
    weather = pd.read_csv(path)
    weather.columns = weather.columns.str.strip()
    weather["date"] = pd.to_datetime(weather["DATE"], errors="coerce").dt.date
    weather["rain_mm"] = (
        pd.to_numeric(weather["neerslag"], errors="coerce").replace(-1, 0) / 10.0
    )
    sunshine_minutes = (
        pd.to_numeric(weather["zonneschijnduur"], errors="coerce")
        .replace(-1, 0)
        .div(10.0)
        .mul(60.0)
    )
    average_temperature_c = pd.to_numeric(
        weather["temperatuur"], errors="coerce"
    ).div(10.0)
    maximum_temperature_c = pd.to_numeric(
        weather["maxtemperatuur"], errors="coerce"
    ).div(10.0)
    weather["energy_proxy"] = calculate_energy_proxy(
        sunshine_minutes,
        average_temperature_c,
        maximum_temperature_c,
    )
    return (
        weather[["date", "rain_mm", "energy_proxy"]]
        .dropna(subset=["date"])
        .drop_duplicates(subset="date", keep="last")
        .sort_values("date")
        .reset_index(drop=True)
    )


def load_forecast(as_of: date, db_path: Path = DB_PATH) -> tuple[date | None, pd.DataFrame]:
    """Lees de nieuwste weersverwachting die op de uitvoerdatum beschikbaar is."""
    with closing(sqlite3.connect(db_path)) as connection:
        forecast_dates = pd.read_sql_query(
            """
            SELECT DISTINCT forecast_date
            FROM TenDayForecast
            WHERE forecast_date <= ?
            ORDER BY forecast_date DESC
            """,
            connection,
            params=(as_of.isoformat(),),
        )
        if forecast_dates.empty:
            return None, pd.DataFrame()
        forecast_date = pd.to_datetime(forecast_dates.iloc[0]["forecast_date"]).date()
        forecast = pd.read_sql_query(
            """
            SELECT
                forecast_for_date,
                precipitation_mm,
                precipitation_perc,
                sunshine_minutes,
                avg_temp,
                max_temp
            FROM TenDayForecast
            WHERE forecast_date = ?
            ORDER BY forecast_for_date
            """,
            connection,
            params=(forecast_date.isoformat(),),
        )

    forecast["date"] = pd.to_datetime(
        forecast["forecast_for_date"], errors="coerce"
    ).dt.date
    forecast["rain_mm"] = pd.to_numeric(
        forecast["precipitation_mm"], errors="coerce"
    ).fillna(0.0)
    forecast["rain_probability"] = pd.to_numeric(
        forecast["precipitation_perc"], errors="coerce"
    ).fillna(100.0).clip(0.0, 100.0)
    forecast["energy_proxy"] = calculate_energy_proxy(
        pd.to_numeric(forecast["sunshine_minutes"], errors="coerce"),
        pd.to_numeric(forecast["avg_temp"], errors="coerce"),
        pd.to_numeric(forecast["max_temp"], errors="coerce"),
    )
    return forecast_date, forecast[
        ["date", "rain_mm", "rain_probability", "energy_proxy"]
    ].dropna(subset=["date"])


def probe_key(device_id: int, probe_number: int, depth_cm: int | float) -> str:
    """Maak de blijvende identiteit van een probe inclusief exacte diepte."""
    return f"{int(device_id)}:{int(probe_number)}:{int(depth_cm)}"


def calculate_critical_permittivity(
    values: pd.Series,
    quantile: float,
) -> float | None:
    """Bereken een probe-eigen lage permittiviteitsgrens uit trainingsdata."""
    numeric_values = pd.to_numeric(values, errors="coerce").dropna()
    if len(numeric_values) < MINIMUM_TRAINING_DAYS:
        return None
    threshold = float(numeric_values.quantile(quantile))
    if not np.isfinite(threshold):
        return None
    return threshold


def calculate_archived_statistics(
    probe_data: pd.DataFrame,
    weather: pd.DataFrame,
    rain_threshold_mm: float,
    cutoff_date: date | None = None,
) -> tuple[float, float] | None:
    """Bereken regenstijging en uitdroging volgens het gearchiveerde model."""
    history = probe_data.copy()
    if cutoff_date is not None:
        history = history[history["date"] <= cutoff_date]
    combined = (
        history[["date", "relative_permittivity"]]
        .merge(weather[["date", "rain_mm"]], on="date", how="left")
        .dropna(subset=["relative_permittivity", "rain_mm"])
        .sort_values("date")
        .reset_index(drop=True)
    )
    if combined.empty:
        return None
    combined["major_rain_event"] = combined["rain_mm"] > rain_threshold_mm
    combined["delta_2_days"] = (
        combined["relative_permittivity"].shift(-2)
        - combined["relative_permittivity"]
    )
    rain_rises = combined.loc[combined["major_rain_event"], "delta_2_days"].dropna()
    if rain_rises.empty:
        return None
    average_rise = float(rain_rises.mean() / 2.0)

    slopes = []
    rain_indices = combined.index[combined["major_rain_event"]].tolist()
    for index in range(len(rain_indices) - 1):
        rain_index = rain_indices[index]
        next_rain_index = rain_indices[index + 1]
        peak_index = rain_index + 2
        if peak_index >= len(combined) or next_rain_index <= peak_index:
            continue
        days = (
            combined.loc[next_rain_index, "date"]
            - combined.loc[peak_index, "date"]
        ).days
        if days <= 5:
            continue
        slopes.append(
            (
                combined.loc[next_rain_index, "relative_permittivity"]
                - combined.loc[peak_index, "relative_permittivity"]
            ) / days
        )
    if not slopes:
        return None
    return average_rise, float(np.mean(slopes))


def predict_archived_path(
    current_value: float,
    recent_min: float,
    recent_max: float,
    origin_date: date,
    rain_by_date: dict[date, float],
    parameters: dict[str, float],
    average_rise: float,
    average_slope: float,
) -> np.ndarray:
    """Voer de oorspronkelijke zeven-daagse voorspellingsvergelijking uit."""
    threshold = float(parameters["rain_threshold_mm"])
    strength = float(parameters["rain_response_strength"])
    cap = float(parameters["max_rain_response_multiplier"])
    decay = float(parameters["rain_memory_decay"])
    recent_range = max(float(recent_max) - float(recent_min), 0.05)
    value = float(current_value)
    predictions = []

    for day_number in range(1, FORECAST_HORIZON_DAYS + 1):
        prediction_date = origin_date + timedelta(days=day_number)
        effective_rain = 0.0
        for lag in range(RAIN_MEMORY_DAYS):
            rain = float(rain_by_date.get(prediction_date - timedelta(days=lag), 0.0))
            if rain >= threshold:
                effective_rain += (rain - threshold) * decay ** lag

        normalized_position = (value - recent_min) / recent_range
        if effective_rain >= threshold:
            response = min(cap, effective_rain / threshold * strength)
            dampening = 1.0 - normalized_position
            value += average_rise * response * (0.3 + 0.7 * dampening)
        else:
            value += average_slope * (0.5 + normalized_position)
        predictions.append(value)
    return np.asarray(predictions, dtype=float)


def write_excel_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    """Schrijf leesbare werkbladen met dezelfde sobere opmaak."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)
            worksheet = writer.book[sheet_name[:31]]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(vertical="top")
            for column_number, cells in enumerate(worksheet.columns, start=1):
                width = min(
                    45,
                    max(10, max(len(str(cell.value or "")) for cell in cells) + 2),
                )
                worksheet.column_dimensions[get_column_letter(column_number)].width = width


def probe_visualisation_directory(
    root: Path,
    device_id: int,
    probe_number: int,
    depth_cm: int | float,
) -> Path:
    """Maak een stabiel visualisatiepad voor een sensor-probe-diepte."""
    depth_text = f"{float(depth_cm):g}"
    return (
        root
        / f"sensor_{int(device_id)}"
        / f"probe_{int(probe_number)}_{depth_text}cm"
    )


def write_permittivity_weather_graph(
    path: Path,
    title: str,
    actual: pd.DataFrame,
    weather: pd.DataFrame,
    weather_columns: tuple[str, ...] = ("rain_mm", "energy_proxy"),
    prediction_dates: list[date] | None = None,
    prediction_values: np.ndarray | None = None,
    origin_date: date | None = None,
) -> None:
    """Schrijf permittiviteit met een of meer gekoppelde weerpanelen."""
    panel_definitions = {
        "rain_mm": (
            "Neerslag (mm)",
            "#4C78A8",
            "Permittiviteit en neerslag",
        ),
        "energy_proxy": (
            "Energieproxy (minuten x graden C)",
            "#E39C37",
            "Permittiviteit en energieproxy",
        ),
    }
    if not weather_columns or any(
        column not in panel_definitions for column in weather_columns
    ):
        raise ValueError("Onbekende of ontbrekende weerpanelen")
    if (prediction_dates is None) != (prediction_values is None):
        raise ValueError("Voorspellingsdatums en -waarden moeten samen worden opgegeven")

    path.parent.mkdir(parents=True, exist_ok=True)
    figure, axes = plt.subplots(
        len(weather_columns),
        1,
        figsize=(14, 9 if len(weather_columns) > 1 else 6),
        sharex=True,
        squeeze=False,
    )
    axes = axes[:, 0]
    for axis, column in zip(axes, weather_columns):
        bar_label, bar_color, panel_title = panel_definitions[column]
        bar_axis = axis.twinx()
        bar_axis.bar(
            weather["date"],
            weather[column],
            width=0.8,
            alpha=0.22,
            color=bar_color,
            label=bar_label,
            zorder=1,
        )
        axis.plot(
            actual["date"],
            actual["relative_permittivity"],
            color="#1B4D3E",
            linewidth=2.0,
            label="Gemeten permittiviteit",
            zorder=3,
        )
        if prediction_dates is not None and prediction_values is not None:
            axis.plot(
                prediction_dates,
                prediction_values,
                color="#C44536",
                linewidth=2.0,
                linestyle="--",
                marker="o",
                markersize=3,
                label="Voorspelde permittiviteit",
                zorder=4,
            )
        if origin_date is not None:
            axis.axvline(
                origin_date,
                color="#555555",
                linewidth=1.2,
                linestyle=":",
                label="Voorspellingsstart",
                zorder=2,
            )
        axis.set_ylabel("Relatieve permittiviteit")
        bar_axis.set_ylabel(bar_label)
        axis.set_title(panel_title)
        axis.grid(axis="y", alpha=0.2)
        lines_1, labels_1 = axis.get_legend_handles_labels()
        lines_2, labels_2 = bar_axis.get_legend_handles_labels()
        axis.legend(lines_1 + lines_2, labels_1 + labels_2, loc="upper left")
    axes[-1].set_xlabel("Datum")
    axes[-1].xaxis.set_major_locator(mdates.AutoDateLocator())
    axes[-1].xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m-%d"))
    figure.suptitle(title)
    figure.autofmt_xdate(rotation=35)
    figure.tight_layout(rect=(0, 0, 1, 0.97))
    figure.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(figure)


def create_staging_directory(target: Path) -> Path:
    """Maak een tijdelijke map naast de uiteindelijke visualisatiemap."""
    target.parent.mkdir(parents=True, exist_ok=True)
    return Path(tempfile.mkdtemp(prefix=f".{target.name}_", dir=target.parent))


def replace_generated_directory(staging: Path, target: Path) -> None:
    """Vervang een gegenereerde map en herstel de vorige map bij een fout."""
    backup = target.with_name(f".{target.name}_previous")
    if backup.exists():
        shutil.rmtree(backup)
    if target.exists():
        target.rename(backup)
    try:
        staging.rename(target)
    except Exception:
        if backup.exists() and not target.exists():
            backup.rename(target)
        raise
    if backup.exists():
        shutil.rmtree(backup)


def status_row(inventory_row, run_date: date, status: str) -> dict[str, object]:
    """Maak een volledige operatorregel, ook wanneer voorspellen niet kan."""
    row = {
        "run_date": run_date.isoformat(),
        "device_id": int(inventory_row.device_id),
        "location_name": inventory_row.location_name,
        "probe_number": int(inventory_row.probe_number),
        "depth_cm": inventory_row.depth_cm,
        "status": status,
        "latest_measurement_date": "",
        "latest_relative_permittivity": np.nan,
        "critical_relative_permittivity": np.nan,
        "risk_within_7_days_pct": np.nan,
        "expected_crossing_date": "",
        "retrained_at": "",
        "days_since_retraining": np.nan,
        "training_start": "",
        "training_end": "",
        "training_daily_observations": np.nan,
        "training_raw_readings": np.nan,
        "parameter_training_cases": np.nan,
        "evaluation_cases": np.nan,
        "evaluation_model_mae": np.nan,
        "evaluation_persistence_mae": np.nan,
        "evaluation_beats_persistence": "",
    }
    return row


def run_model(
    as_of: date,
    simulations: int = DEFAULT_SIMULATIONS,
    model_path: Path = MODEL_PATH,
    output_dir: Path = OUTPUT_DIR,
    visualisation_dir: Path = VISUALISATION_DIR,
) -> Path:
    """Maak een operationeel Excel-bestand met een regel per sensor en diepte."""
    print(
        f"Dagelijkse voorspelling gestart | peildatum: {as_of.isoformat()} | "
        f"simulaties per probe: {max(1, simulations)}",
        flush=True,
    )
    with model_path.open(encoding="utf-8") as source:
        artifact = json.load(source)
    if artifact.get("prediction_target") != "relative_permittivity":
        raise ValueError(
            "Het modelbestand is ongeldig. Train het model opnieuw met de huidige "
            "train_permittivity_prediction_model.py."
        )

    usable_sensor_ids = load_usable_sensor_ids()
    inventory = load_active_probe_inventory()
    daily_data = load_daily_sensor_data()
    historical_weather = load_historical_weather()
    forecast_date, forecast = load_forecast(as_of)
    historical_rain = dict(zip(historical_weather["date"], historical_weather["rain_mm"]))
    forecast_by_date = forecast.set_index("date") if not forecast.empty else pd.DataFrame()
    rows = []
    total_probes = len(inventory)
    completed_probes = 0
    predicted_probes = 0

    def record_row(row: dict[str, object], inventory_row) -> None:
        nonlocal completed_probes, predicted_probes
        rows.append(row)
        completed_probes += 1
        if row["status"] == "predicted":
            predicted_probes += 1
        depth = (
            f"{float(inventory_row.depth_cm):g} cm"
            if not pd.isna(inventory_row.depth_cm)
            else "onbekende diepte"
        )
        remaining = total_probes - completed_probes
        remaining_word = "probe" if remaining == 1 else "probes"
        print(
            f"[{completed_probes}/{total_probes}] "
            f"{inventory_row.location_name} | sensor {int(inventory_row.device_id)} | "
            f"probe {int(inventory_row.probe_number)} | {depth} | {row['status']} | "
            f"{remaining} {remaining_word} te gaan",
            flush=True,
        )

    print(
        f"Invoer geladen | {total_probes} actieve "
        f"{'probe' if total_probes == 1 else 'probes'} | "
        f"forecastdatum: {forecast_date}",
        flush=True,
    )
    staging_dir = create_staging_directory(visualisation_dir)
    try:
        for inventory_row in inventory.itertuples(index=False):
            if int(inventory_row.device_id) not in usable_sensor_ids:
                record_row(
                    status_row(inventory_row, as_of, "insufficient_history"),
                    inventory_row,
                )
                continue
            if pd.isna(inventory_row.depth_cm):
                record_row(
                    status_row(inventory_row, as_of, "invalid_depth"),
                    inventory_row,
                )
                continue
            key = probe_key(
                inventory_row.device_id,
                inventory_row.probe_number,
                inventory_row.depth_cm,
            )
            probe_model = artifact["probe_models"].get(key)
            if probe_model is None or probe_model.get("status") != "trained":
                record_row(
                    status_row(inventory_row, as_of, "retraining_required"),
                    inventory_row,
                )
                continue
            trained_at = datetime.fromisoformat(probe_model["trained_at"])
            days_since_retraining = (as_of - trained_at.date()).days

            probe_data = daily_data[
                (daily_data["device_id"] == inventory_row.device_id)
                & (daily_data["probe_number"] == inventory_row.probe_number)
                & (daily_data["depth_cm"] == inventory_row.depth_cm)
                & (daily_data["date"] < as_of)
            ].sort_values("date")
            if probe_data.empty:
                record_row(
                    status_row(inventory_row, as_of, "missing_sensor_data"),
                    inventory_row,
                )
                continue
            latest = probe_data.iloc[-1]
            latest_date = latest["date"]
            if (as_of - latest_date).days > MAX_SENSOR_AGE_DAYS:
                record_row(
                    status_row(inventory_row, as_of, "stale_sensor_data"),
                    inventory_row,
                )
                continue
            if forecast_date is None or (as_of - forecast_date).days > MAX_FORECAST_AGE_DAYS:
                record_row(
                    status_row(inventory_row, as_of, "stale_weather_forecast"),
                    inventory_row,
                )
                continue

            future_dates = [
                latest_date + timedelta(days=day_number)
                for day_number in range(1, FORECAST_HORIZON_DAYS + 1)
            ]
            if not all(day in forecast_by_date.index for day in future_dates):
                record_row(
                    status_row(inventory_row, as_of, "missing_forecast_days"),
                    inventory_row,
                )
                continue

            recent = probe_data[probe_data["date"] >= latest_date - timedelta(days=30)]
            rain_by_date = dict(historical_rain)
            rain_by_date.update(
                forecast_by_date.loc[future_dates, "rain_mm"].astype(float).to_dict()
            )
            point_prediction = predict_archived_path(
                float(latest["relative_permittivity"]),
                float(recent["relative_permittivity"].min()),
                float(recent["relative_permittivity"].max()),
                latest_date,
                rain_by_date,
                probe_model["parameters"],
                float(probe_model["average_rise_per_day"]),
                float(probe_model["average_slope_per_day"]),
            )
            critical_permittivity = float(
                probe_model["critical_relative_permittivity"]
            )

            random_generator = np.random.default_rng(
                as_of.toordinal()
                + int(inventory_row.device_id) * 10
                + int(inventory_row.probe_number)
            )
            residuals = np.asarray(
                probe_model.get("evaluation_residuals_relative_permittivity", []),
                dtype=float,
            )
            simulated_paths = []
            for _ in range(max(1, simulations)):
                sampled_rain = dict(rain_by_date)
                rain_occurs = random_generator.random(FORECAST_HORIZON_DAYS) < (
                    forecast_by_date.loc[
                        future_dates, "rain_probability"
                    ].to_numpy(float)
                    / 100.0
                )
                for prediction_date, occurs in zip(future_dates, rain_occurs):
                    if not occurs:
                        sampled_rain[prediction_date] = 0.0
                simulated = predict_archived_path(
                    float(latest["relative_permittivity"]),
                    float(recent["relative_permittivity"].min()),
                    float(recent["relative_permittivity"].max()),
                    latest_date,
                    sampled_rain,
                    probe_model["parameters"],
                    float(probe_model["average_rise_per_day"]),
                    float(probe_model["average_slope_per_day"]),
                )
                if residuals.ndim == 2 and len(residuals):
                    simulated = (
                        simulated
                        + residuals[random_generator.integers(0, len(residuals))]
                    )
                simulated_paths.append(simulated)
            simulated_paths = np.asarray(simulated_paths)
            crossing_mask = simulated_paths <= critical_permittivity
            risk = float(np.mean(np.any(crossing_mask, axis=1)) * 100.0)
            crossing_days = [
                int(np.argmax(path_crosses)) + 1
                for path_crosses in crossing_mask
                if np.any(path_crosses)
            ]
            expected_crossing_date = (
                latest_date
                + timedelta(days=int(round(float(np.median(crossing_days)))))
                if crossing_days
                else None
            )

            row = status_row(inventory_row, as_of, "predicted")
            row.update(
                {
                    "latest_measurement_date": latest_date.isoformat(),
                    "latest_relative_permittivity": float(
                        latest["relative_permittivity"]
                    ),
                    "critical_relative_permittivity": critical_permittivity,
                    "risk_within_7_days_pct": round(risk, 1),
                    "expected_crossing_date": (
                        expected_crossing_date.isoformat()
                        if expected_crossing_date
                        else ""
                    ),
                    "retrained_at": probe_model["trained_at"],
                    "days_since_retraining": days_since_retraining,
                    "training_start": probe_model["training_start"],
                    "training_end": probe_model["training_end"],
                    "training_daily_observations": probe_model[
                        "training_daily_observations"
                    ],
                    "training_raw_readings": probe_model["training_raw_readings"],
                    "parameter_training_cases": probe_model[
                        "parameter_training_cases"
                    ],
                    "evaluation_cases": probe_model["evaluation"]["cases"],
                    "evaluation_model_mae": probe_model["evaluation"]["model_mae"],
                    "evaluation_persistence_mae": probe_model["evaluation"][
                        "persistence_mae"
                    ],
                    "evaluation_beats_persistence": probe_model["evaluation"][
                        "beats_persistence"
                    ],
                }
            )
            graph_weather = pd.concat(
                [
                    historical_weather[
                        (historical_weather["date"] >= latest_date - timedelta(days=30))
                        & (historical_weather["date"] <= latest_date)
                    ],
                    forecast_by_date.loc[
                        future_dates, ["rain_mm", "energy_proxy"]
                    ].reset_index(),
                ],
                ignore_index=True,
            ).sort_values("date")
            graph_path = probe_visualisation_directory(
                staging_dir,
                inventory_row.device_id,
                inventory_row.probe_number,
                inventory_row.depth_cm,
            ) / "prediction.png"
            write_permittivity_weather_graph(
                path=graph_path,
                title=(
                    f"{inventory_row.location_name} - sensor {int(inventory_row.device_id)}, "
                    f"probe {int(inventory_row.probe_number)}, "
                    f"{float(inventory_row.depth_cm):g} cm"
                ),
                actual=recent[["date", "relative_permittivity"]],
                weather=graph_weather,
                prediction_dates=[latest_date, *future_dates],
                prediction_values=np.concatenate(
                    ([float(latest["relative_permittivity"])], point_prediction)
                ),
                origin_date=latest_date,
            )
            record_row(row, inventory_row)

        output = pd.DataFrame(rows)
        output["_location_missing"] = output["location_name"].isna()
        output["_location_sort"] = (
            output["location_name"].fillna("").astype(str).str.casefold()
        )
        output = (
            output.sort_values(
                [
                    "_location_missing",
                    "_location_sort",
                    "depth_cm",
                    "device_id",
                    "probe_number",
                ],
                ascending=True,
                na_position="last",
            )
            .drop(columns=["_location_missing", "_location_sort"])
            .reset_index(drop=True)
        )
    except Exception:
        shutil.rmtree(staging_dir, ignore_errors=True)
        raise
    output_path = output_dir / f"prediction_{as_of.isoformat()}.xlsx"
    run_information = pd.DataFrame(
        [
            {"item": "run_date", "value": as_of.isoformat()},
            {"item": "model_file_created_at", "value": artifact["created_at"]},
            {"item": "training_method", "value": artifact["training_method"]},
            {"item": "weather_forecast_date", "value": forecast_date},
            {"item": "model_file", "value": str(model_path)},
        ]
    )
    try:
        write_excel_workbook(
            output_path,
            {"Predictions": output, "Run information": run_information},
        )
        replace_generated_directory(staging_dir, visualisation_dir)
    except Exception:
        shutil.rmtree(staging_dir, ignore_errors=True)
        raise
    print(
        f"Dagelijkse voorspelling voltooid | {completed_probes} "
        f"{'probe' if completed_probes == 1 else 'probes'} verwerkt | "
        f"{predicted_probes} voorspellingen en grafieken | Excel: {output_path} | "
        f"grafieken: {visualisation_dir}",
        flush=True,
    )
    return output_path


def parse_date(value: str) -> date:
    return date.fromisoformat(value)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", nargs="?", choices=["run"], default="run")
    parser.add_argument("--as-of", type=parse_date, default=date.today())
    parser.add_argument("--simulations", type=int, default=DEFAULT_SIMULATIONS)
    parser.add_argument("--model-path", type=Path, default=MODEL_PATH)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument(
        "--visualisation-dir",
        type=Path,
        default=VISUALISATION_DIR,
    )
    return parser


def main() -> None:
    arguments = build_parser().parse_args()
    try:
        run_model(
            arguments.as_of,
            arguments.simulations,
            arguments.model_path,
            arguments.output_dir,
            arguments.visualisation_dir,
        )
    except Exception as error:
        print(f"Dagelijkse voorspelling gestopt met fout | {error}", flush=True)
        raise


if __name__ == "__main__":
    main()
